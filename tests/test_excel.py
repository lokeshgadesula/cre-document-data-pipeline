from openpyxl import Workbook
from cre_pipeline.excel_normalizer import expand_merged_cells, hierarchical_headers

def test_merged_headers_expand_and_flatten():
    wb=Workbook(); ws=wb.active
    ws.merge_cells("A1:B1"); ws["A1"]="Year 2024"; ws["A2"]="Jan"; ws["B2"]="Feb"
    expand_merged_cells(ws)
    assert ws["B1"].value == "Year 2024"
    assert hierarchical_headers(ws,2) == ["year_2024__jan","year_2024__feb"]
